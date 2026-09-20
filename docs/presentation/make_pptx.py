"""Generate docs/presentation/enjoy-pitch.pptx — 2-minute hackathon deck, v2.

Run:  Backend/.venv/Scripts/python.exe docs/presentation/make_pptx.py
Re-run any time after editing; the .pptx is overwritten (close PowerPoint first!).

v2: owner's retyped content (slides 2-3 verbatim), emoji bullets, short
fragments (no full sentences), chip/pill layout. Locked palette, NO red.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Locked palette (frontend/lib/theme.ts)
CREAM = RGBColor(0xED, 0xED, 0xDD)
NAVY = RGBColor(0x1A, 0x20, 0x3B)
TEAL = RGBColor(0x2E, 0x8C, 0x8C)
DEEP_TEAL = RGBColor(0x0F, 0x4C, 0x5C)
MIST = RGBColor(0xD8, 0xE4, 0xE0)
SOFT_AMBER = RGBColor(0xE8, 0xB4, 0x6A)
WHITE = RGBColor(0xF7, 0xF7, 0xF0)
DARK_CHIP = RGBColor(0x24, 0x2C, 0x4E)
LIGHT_CHIP = RGBColor(0xFF, 0xFF, 0xF8)

FONT = "Segoe UI"
EMOJI_FONT = "Segoe UI Emoji"

SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]

page_no = 0


def slide(bg):
    global page_no
    page_no += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s, page_no


def box(s, left, top, width, height, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.5)
    sh.shadow.inherit = False
    return sh


def txt(s, left, top, width, height, lines, *, align=PP_ALIGN.CENTER, anchor=None):
    """lines: list of (text, size, color, bold) or (text, size, color, bold, font)."""
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, spec in enumerate(lines):
        t, size, color, bold = spec[0], spec[1], spec[2], spec[3]
        font = spec[4] if len(spec) > 4 else FONT
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = t
        run.font.name = font
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return tb


def footer(s, n, dark_bg):
    col = MIST if dark_bg else TEAL
    box(s, 0.55, 7.08, 0.5, 0.045, fill=col)
    txt(s, 0.55, 6.68, 3.0, 0.4, [("enjoy ⚓", 12, col, True)], align=PP_ALIGN.LEFT)
    txt(s, 12.2, 6.68, 0.6, 0.4, [(f"{n:02d}", 12, col, True)], align=PP_ALIGN.RIGHT)


def chip(s, left, top, width, height, fill, lines, *, align=PP_ALIGN.LEFT):
    box(s, left, top, width, height, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    txt(s, left + 0.35, top + 0.12, width - 0.7, height - 0.24, lines, align=align)


def kicker(s, text_, color):
    txt(s, 0.9, 0.55, 6.0, 0.5, [(text_, 16, color, True)], align=PP_ALIGN.LEFT)


def title(s, text_, color, size=44, top=1.0):
    txt(s, 0.85, top, 11.6, 1.3, [(text_, size, color, True)], align=PP_ALIGN.LEFT)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ── S1 · Title ──────────────────────────────────────────────────────────────
s, n = slide(NAVY)
box(s, 0, 5.75, SW, 0.1, fill=TEAL)
txt(s, 1, 0.9, 11.3, 1.0, [("🌊 ⛵ 🏮", 40, WHITE, False, EMOJI_FONT)])
txt(s, 1, 1.9, 11.3, 2.6, [
    ("enjoy", 100, WHITE, True),
    ("the Harbour of Learning", 34, SOFT_AMBER, False),
])
txt(s, 1, 4.7, 11.3, 0.8, [("a cozy ocean world that automates studying", 24, CREAM, False)])
footer(s, n, dark_bg=True)
notes(s, "0:00-0:10 — 'This is enjoy — a calm ocean world that does the busywork of studying for you.'")

# ── S2 · Problem (owner's text) ─────────────────────────────────────────────
s, n = slide(CREAM)
kicker(s, "THE PROBLEM 😰", DEEP_TEAL)
title(s, "Students are stressed out 🎓💦", NAVY)
rows = [
    ("📚", "current edtech  →  adds MORE burden"),
    ("⏳", "time wasted on tools  →  better spent learning yourself"),
    ("🤖", "or a chatbot finishes the work  →  nothing is learned"),
]
y = 2.55
for emoji, line in rows:
    chip(s, 1.1, y, 11.1, 0.95, LIGHT_CHIP, [
        (f"{emoji}  {line}", 24, NAVY, False),
    ])
    y += 1.2
footer(s, n, dark_bg=False)
notes(s, "0:10-0:25 — owner's framing: stress, burden, wasted time, chatbot shortcut. Land on: 'none of that is learning.'")

# ── S3 · What we offer (owner's text) ───────────────────────────────────────
s, n = slide(DEEP_TEAL)
kicker(s, "WHAT WE OFFER ✅", SOFT_AMBER)
title(s, "Calm automation, real learning 🌊", WHITE)
rows = [
    ("⏱️", "time-saving AI agents  →  boring tasks gone (notes, planning)"),
    ("🧩", "big tasks  →  broken down  →  executive function boost"),
    ("🧘", "psychology-backed calm  →  ❌ dashboards  ❌ red deadlines"),
]
y = 2.55
for emoji, line in rows:
    chip(s, 1.1, y, 11.1, 0.95, DARK_CHIP, [
        (f"{emoji}  {line}", 24, WHITE, False),
    ])
    y += 1.2
footer(s, n, dark_bg=True)
notes(s, "0:25-0:40 — three offers, one breath each. Then: 'let me show you.'")

# ── S4 · Demo A: chat -> plan ───────────────────────────────────────────────
s, n = slide(NAVY)
kicker(s, "DEMO 1/3 💬", SOFT_AMBER)
title(s, "One sentence  →  a plan 🗺️", WHITE)
chip(s, 1.1, 2.7, 11.1, 1.15, DARK_CHIP, [
    ('"I have a thermodynamics exam Friday"', 26, SOFT_AMBER, True),
    ("→  roadmap   →  tasks   →  boats on the water ⛵", 22, CREAM, False),
])
txt(s, 1, 4.9, 11.3, 0.9, [
    ("▶  LIVE: type in chat  →  tour offer  →  boats spawn", 20, TEAL, True),
])
footer(s, n, dark_bg=True)
notes(s, "0:40-1:05 — CUT TO GAMEPLAY A. 'The orchestrator routes it, the planner templates it, the world shows it.'")

# ── S5 · Demo B: review loop ────────────────────────────────────────────────
s, n = slide(TEAL)
kicker(s, "DEMO 2/3 ⛵", NAVY)
title(s, "Flashcards  =  a short sail 🃏", WHITE)
chip(s, 1.1, 2.6, 11.1, 1.7, DARK_CHIP, [
    ("board  →  ❓ question  →  👀 reveal  →  4️⃣ grade boats", 26, WHITE, True),
    ("deck done  →  docks at the lamp  →  glows brighter 🏮✨", 22, CREAM, False),
])
txt(s, 1, 4.9, 11.3, 0.9, [
    ("▶  LIVE: 3-card review, keyboard only, dock + lamp glow", 20, NAVY, True),
])
footer(s, n, dark_bg=True)
notes(s, "1:05-1:30 — CUT TO GAMEPLAY B. 'The AI never answers for you — it removes everything else.'")

# ── S6 · Demo C: atlas ──────────────────────────────────────────────────────
s, n = slide(NAVY)
kicker(s, "DEMO 3/3 🫧", SOFT_AMBER)
title(s, "Knowledge grows under the sea 🪸", WHITE)
chip(s, 1.1, 2.7, 11.1, 1.15, DARK_CHIP, [
    ("dive down 🤿  →  every concept joins a living atlas", 24, WHITE, False),
    ("your portfolio isn't a chart  →  it's a place", 22, CREAM, False),
])
txt(s, 1, 4.9, 11.3, 0.9, [
    ("▶  LIVE: dive transition  →  atlas  →  click a node", 20, TEAL, True),
])
footer(s, n, dark_bg=True)
notes(s, "1:30-1:45 — CUT TO GAMEPLAY C: dive, pan, one node drawer, surface.")

# ── S7 · Tech stack (one page: 3 whys on top, full list below) ─────────────
s, n = slide(CREAM)
kicker(s, "TECH STACK ⚙️", DEEP_TEAL)
title(s, "Why this stack 🧰", NAVY, size=38, top=0.95)

# Three talking points the owner presents verbally
whys = [
    ("🤖", "Multi-agent system", "specialised agents  →  higher quality"),
    ("🦜🔗", "LangChain", "guides the AI  →  controlled, reliable output"),
    ("🕸️", "Three.js", "powers the 3D game world in the browser"),
]
x = 1.1
for emoji, name, why in whys:
    chip(s, x, 2.05, 3.57, 1.5, LIGHT_CHIP, [
        (f"{emoji}  {name}", 22, NAVY, True),
        (why, 16, DEEP_TEAL, False),
    ])
    x += 3.77

# Full stack list, grouped
box(s, 1.1, 3.85, 11.1, 0.5, fill=DEEP_TEAL)
txt(s, 1.45, 3.92, 10.4, 0.4, [("THE FULL STACK", 16, WHITE, True)], align=PP_ALIGN.LEFT)
cols = [
    ("FRONTEND 🖥️", [
        "Next.js 16  ·  React 19  ·  TypeScript",
        "Three.js  +  react-three-fiber",
        "Zustand  ·  Tailwind CSS",
    ]),
    ("BACKEND 🧠", [
        "Python 3.12  ·  FastAPI  ·  Pydantic",
        "Multi-agent: orchestrator · planner · flashcards",
        "LangChain / LangGraph  ·  SSE streaming",
    ]),
    ("DATA & INFRA 🗄️", [
        "PostgreSQL 16  ·  SQLAlchemy async",
        "Alembic migrations  ·  Docker Compose",
        "event-sourced world state  →  never drifts",
    ]),
]
x = 1.1
for head, items in cols:
    box(s, x, 4.35, 3.57, 2.15, fill=MIST)
    txt(s, x + 0.3, 4.5, 3.0, 0.4, [(head, 15, DEEP_TEAL, True)], align=PP_ALIGN.LEFT)
    txt(s, x + 0.3, 4.95, 3.05, 1.5,
        [(i, 13.5, NAVY, False) for i in items], align=PP_ALIGN.LEFT)
    x += 3.77
footer(s, n, dark_bg=False)
notes(s, "1:45-1:55 — talk ONLY the top 3: multi-agent for quality, LangChain to guide the AI, Three.js for the world. Point at the full list, don't read it.")

# ── S8 · Standing on open source (credits, per agent — opensource/README.md) ─
s, n = slide(DEEP_TEAL)
kicker(s, "STANDING ON OPEN SOURCE 💙", SOFT_AMBER)
title(s, "Built with open source, agent by agent 🌊", WHITE, size=36, top=0.95)

oss = [
    ("🛳️ Scheduler agent",
     "LangGraph plan-and-execute  ·  obsidian-magic-tasks  ·  study-revision-planner"),
    ("⛵ Flashcard agent",
     "anki-llm-flashcard-generator (forked)  ·  py-fsrs  ·  genanki"),
    ("🌊 Knowledge graph",
     "langchain-experimental graph transformer  ·  nano-graphrag  ·  KeyBERT"),
    ("🧭 Orchestrator",
     "fastapi-langgraph-template  ·  langgraph-swarm handoffs"),
    ("🎨 World rendering",
     "stylized-water shaders  ·  camera-controls  ·  underwater-atlas engine (three.js + GLSL)"),
    ("🧱 3D assets",
     "Kenney CC0 boats  ·  custom Blender fishboat + lamp buoy"),
]
for i, (head, line) in enumerate(oss):
    x = 1.1 + (i % 3) * 3.77
    y = 2.05 + (i // 3) * 2.1
    chip(s, x, y, 3.57, 1.85, DARK_CHIP, [
        (head, 17, SOFT_AMBER, True),
        (line, 13.5, CREAM, False),
    ])
txt(s, 1, 6.25, 11.3, 0.5, [
    ("full credit list  →  opensource/README.md 📜", 15, MIST, False),
])
footer(s, n, dark_bg=True)
notes(s, "Credit slide — not spoken in the 2:00. If asked about tech: 'each agent stands on proven open source — the list is in our repo.'")

# ── S9 · Closing ────────────────────────────────────────────────────────────
s, n = slide(NAVY)
txt(s, 1, 1.3, 11.3, 1.0, [("🏮", 54, WHITE, False, EMOJI_FONT)])
txt(s, 1, 2.4, 11.3, 1.2, [
    ("The lamp records victories only.", 46, WHITE, True),
])
box(s, 4.9, 3.85, 3.53, 0.06, fill=SOFT_AMBER)
txt(s, 1, 4.15, 11.3, 0.8, [
    ("enjoy — come back to a harbour, never a bill. ⚓", 26, SOFT_AMBER, False),
])
txt(s, 1, 5.5, 11.3, 0.8, [("thank you 🙏", 22, CREAM, False)])
footer(s, n, dark_bg=True)
notes(s, "1:55-2:00 — slow: 'A returning student is welcomed, never billed for absence. Thank you.'")

prs.save("docs/presentation/enjoy-pitch.pptx")
print(f"OK — {page_no} slides")
